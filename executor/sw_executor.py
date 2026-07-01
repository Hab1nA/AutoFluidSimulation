"""
===============================================================================
SolidWorks COM 自动化执行器 (SW Executor)

负责通过 win32com 驱动 SolidWorks 完成：
- SW 进程连接/启动（三层降级策略）
- 使用模型内已链接设计表提供的构型
- 所有构型 STEP 文件导出
- COM 资源清理

从 engine/task_runner.py 中提取，职责独立后可单独测试和维护。
===============================================================================
"""
from __future__ import annotations

import os
import subprocess
import time
import gc
import threading
from collections.abc import Iterator
from contextlib import contextmanager
from typing import Any, TYPE_CHECKING

if TYPE_CHECKING:
    from engine.scheduler.control import PipelineControl
    from engine.state_manager import StateManager

from engine.config import (
    LOCAL_PATHS, ENGINE_CONFIG,
    OPERATION_TIMEOUTS,
    STATUS_COMPLETED, STATUS_ERROR, get_step_filename,
)
from utils.logger import setup_logger

logger = setup_logger(__name__)


class SWExecutor:
    """SolidWorks COM 自动化执行器。

    封装 SolidWorks 的连接、模型构型切换、STEP 导出等所有 COM 操作。
    """

    # swDocumentTypes_e
    _SW_DOC_PART = 1
    _SW_DOC_ASSEMBLY = 2
    # swOpenDocOptions_e (用于 OpenDoc6)
    _SW_OPEN_SILENT = 1       # 静默：抑制警告对话框
    _SW_OPEN_READONLY = 2     # 只读
    # swRunMacroOption_e (用于 RunMacro2)
    _SW_RUN_MACRO_DEFAULT = 0
    _SW_RUN_MACRO_UNLOAD_AFTER = 1
    # swSaveAsVersion_e / swSaveAsOptions_e (用于直接 COM 导出 STEP，替代宏文件)
    _SW_SAVE_AS_CURRENT_VERSION = 0   # swSaveAsCurrentVersion
    _SW_SAVE_AS_OPTIONS_SILENT = 1    # swSaveAsOptions_Silent

    def __init__(self, state_manager: StateManager):
        """初始化 SW 执行器。

        Args:
            state_manager: StateManager 实例
        """
        self.state = state_manager
        self._paused_event: threading.Event | None = None
        self._stopped_event: threading.Event | None = None
        self._pipeline_control: PipelineControl | None = None
        self._cached_sw_app: Any = None
        self._cached_doc: Any = None
        self._com_initialized: bool = False
        self._cleanup_lock = threading.RLock()
        self._first_cleanup_done = False
        self._final_cleanup_done = False
        self.last_error = ""
        self._last_open_error: Exception | None = None

    def set_control_events(
        self,
        paused_event: threading.Event,
        stopped_event: threading.Event,
    ) -> None:
        """注入调度器的暂停/停止事件，使逐构型循环可响应 pause/stop。

        Args:
            paused_event: 调度器的 _paused 事件
            stopped_event: 调度器的 _stopped 事件
        """
        self._paused_event = paused_event
        self._stopped_event = stopped_event

    def set_pipeline_control(self, pipeline_control: PipelineControl) -> None:
        """注入统一控制层，用于锁定单构型 COM 操作窗口。"""
        self._pipeline_control = pipeline_control

    @contextmanager
    def _external_start(self) -> Iterator[bool]:
        """锁定一个 SW 外部操作窗口。"""
        if self._pipeline_control is not None:
            with self._pipeline_control.external_start() as allowed:
                yield allowed
            return
        paused = self._paused_event is not None and self._paused_event.is_set()
        stopped = self._stopped_event is not None and self._stopped_event.is_set()
        yield not paused and not stopped

    # ------------------------------------------------------------------
    # 单构型导出（供 RetryManager 逐构型调用）
    # ------------------------------------------------------------------

    def export_sw_per_config(self, config_name: int) -> bool:
        """在统一 gate 下执行单构型 SW 导出。"""
        self.last_error = ""
        with self._external_start() as allowed:
            if not allowed:
                return self._fail(f"[SW] 构型{config_name}: 暂停或停止状态下跳过导出")
            return self._export_sw_per_config_admitted(config_name)

    def _fail(self, message: str) -> bool:
        self.last_error = message
        logger.error(message)
        return False

    def _export_sw_per_config_admitted(self, config_name: int) -> bool:
        """导出单个构型的 STEP 文件（供 RetryManager 调用）。

        执行流程：
        1. 首次调用时建立 SW 连接、打开已链接设计表的模型（缓存在实例属性）
        2. 切换到目标构型 → 重建 → SaveAs STEP
        3. 所有构型完成后由调用方调用 _disconnect_sw_cached() 清理

        由 RetryManager 逐构型调用，
          状态管理（Running/Retrying/Error）完全由 RetryManager 负责

        Args:
            config_name: 构型编号

        Returns:
            True 表示该构型 STEP 文件导出成功
        """
        # ★ CoInitialize 仅在首次调用时执行（同一线程内幂等）
        if not self._com_initialized:
            try:
                import pythoncom
                pythoncom.CoInitialize()
                self._com_initialized = True
            except ImportError:
                return self._fail("[SW] pywin32 未安装，无法初始化 COM")
            except Exception:
                pass  # 已初始化

        step_dir = LOCAL_PATHS.get("step_dir", "")
        sw_model = LOCAL_PATHS["sw_model"]
        doc_type = self._guess_sw_doc_type(sw_model)

        # ---- 首次调用：建立 SW 连接 ----
        if self._cached_sw_app is None:
            sw_app = None
            doc = None
            cache_ready = False
            try:
                sw_app = self._connect_sw()
                if sw_app is None:
                    return self._fail(f"[SW] 构型{config_name}: 无法连接 SolidWorks")
                try:
                    sw_app.Visible = bool(ENGINE_CONFIG.get("sw_visible", True))
                except Exception:
                    pass
                doc = self._open_sw_model(sw_app, sw_model, doc_type)
                if doc is None:
                    recovered, recovered_sw_app, recovered_doc = (
                        self._recover_after_open_rpc_failure(
                            sw_model,
                            doc_type,
                        )
                    )
                    if recovered:
                        sw_app = recovered_sw_app
                        doc = recovered_doc
                if doc is None:
                    return self._fail(f"[SW] 构型{config_name}: 无法打开模型")
                logger.info(
                    "[SW] 模型已打开；使用模型内已链接设计表，"
                    "不主动导入 Excel 或批量写入参数"
                )
                self._cached_sw_app = sw_app
                self._cached_doc = doc
                cache_ready = True
                logger.info("[SW] COM 连接已建立（缓存供后续构型复用）")
            except Exception as e:
                logger.error(
                    f"[SW] 构型{config_name}: SW 连接失败 "
                    f"({type(e).__name__}: {e})", exc_info=True
                )
                self.last_error = f"[SW] 构型{config_name}: SW 连接失败 ({type(e).__name__}: {e})"
                return False
            finally:
                if not cache_ready:
                    if sw_app is not None or doc is not None:
                        try:
                            self._disconnect_sw(sw_app, doc, sw_model)
                        except Exception as e:
                            logger.debug(f"[SW-Cleanup] 初始化失败后清理异常: {e}")
                        finally:
                            self._com_initialized = False
                    else:
                        self._uninitialize_com_if_needed()

        sw_app = self._cached_sw_app
        doc = self._cached_doc

        # ---- 校验 STEP 目录 ----
        if not step_dir:
            return self._fail("[SW] 未配置 STEP 输出目录 (step_dir)")
        try:
            os.makedirs(step_dir, exist_ok=True)
        except OSError as e:
            return self._fail(f"[SW] 无法创建 STEP 输出目录: {step_dir}: {e}")

        # ---- 检查文件是否已存在（断点续传 / 之前批次已成功） ----
        filename = get_step_filename("sw", config_name)
        if not filename:
            return self._fail(f"[SW] 构型{config_name}: 无法生成 STEP 文件名")
        filepath = os.path.join(step_dir, filename)
        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            logger.info(
                f"[SW] 构型{config_name}: STEP 文件已存在，跳过导出"
            )
            return True

        # ---- 暂停/停止检查 ----
        if self._paused_event is not None and self._paused_event.is_set():
            self.last_error = f"[SW] 构型{config_name}: 暂停标志已置位，中止导出"
            logger.info(self.last_error)
            return False
        if self._stopped_event is not None and self._stopped_event.is_set():
            self.last_error = f"[SW] 构型{config_name}: 停止标志已置位，中止导出"
            logger.info(self.last_error)
            return False

        # ---- 切换构型 → 重建 → 导出 ----
        import win32com.client

        logger.info(f"[SW] 构型{config_name}: 开始导出 STEP...")
        cn_str = str(config_name)
        design_table_snapshot = self._capture_design_table_snapshot()

        try:
            doc.ShowConfiguration2(cn_str)
        except Exception as e:
            self.disconnect_sw_cached()
            return self._fail(
                f"[SW] 构型{config_name}: ShowConfiguration2 失败 "
                f"({type(e).__name__}: {e})"
            )

        # 重建
        ext = doc.Extension
        rebuild_ok = False
        try:
            ext.Rebuild(0)
            rebuild_ok = True
        except TypeError:
            try:
                rebuild_fn = ext.Rebuild
                if callable(rebuild_fn):
                    rebuild_fn(0)
                    rebuild_ok = True
            except Exception:
                pass
        except Exception:
            pass
        if not rebuild_ok:
            try:
                doc.Rebuild(0)
                rebuild_ok = True
            except Exception:
                pass

        # SaveAs STEP
        import pythoncom
        try:
            save_errors = win32com.client.VARIANT(
                pythoncom.VT_BYREF | pythoncom.VT_I4, 0
            )
            save_warnings = win32com.client.VARIANT(
                pythoncom.VT_BYREF | pythoncom.VT_I4, 0
            )
            export_data = win32com.client.VARIANT(
                pythoncom.VT_DISPATCH, None
            )
            status = doc.Extension.SaveAs(
                filepath,
                self._SW_SAVE_AS_CURRENT_VERSION,
                self._SW_SAVE_AS_OPTIONS_SILENT,
                export_data,
                save_errors,
                save_warnings,
            )
            if status and os.path.exists(filepath):
                logger.info(
                    f"[SW] 构型{config_name}: {filename} 导出成功 "
                    f"(Errors={save_errors.value}, Warnings={save_warnings.value})"
                )
                return True
            else:
                self.last_error = (
                    f"[SW] 构型{config_name}: SaveAs 返回 {status}，"
                    f"文件存在={os.path.exists(filepath)} "
                    f"(Errors={save_errors.value})"
                )
                logger.warning(self.last_error)
                return False
        except Exception as e:
            self.last_error = (
                f"[SW] 构型{config_name}: SaveAs 异常 "
                f"({type(e).__name__}: {e})"
            )
            logger.error(self.last_error)
            # ★ 清理缓存的 COM 连接，下次重试时重新建立
            self.disconnect_sw_cached()
            if self._is_com_rpc_failure(e):
                logger.warning(
                    "[SW-COM] SaveAs 检测到 COM/RPC 失败，"
                    "将终止残留 SolidWorks 进程以便下次重试重新启动"
                )
                self._terminate_sw_processes()
            return False
        finally:
            if not self._restore_design_table_snapshot(design_table_snapshot):
                return False

    def disconnect_sw_cached(self) -> None:
        """清理通过 export_sw_per_config 缓存的 SW 连接。"""
        sw_app = self._cached_sw_app
        doc = self._cached_doc
        self._cached_sw_app = None
        self._cached_doc = None

        if sw_app is None and doc is None:
            self._uninitialize_com_if_needed()
            return

        sw_model = LOCAL_PATHS.get("sw_model", "")
        try:
            self._disconnect_sw(sw_app, doc, sw_model)
        except Exception as e:
            logger.debug(f"[SW-Cleanup] 清理缓存连接异常: {e}")
        finally:
            self._com_initialized = False

    # ------------------------------------------------------------------
    # 全量清理（首次/末次）
    # ------------------------------------------------------------------

    def shutdown_all(self) -> None:
        """全量清理 SolidWorks 进程并重置清理标志。"""
        with self._cleanup_lock:
            self._shutdown_all_internal()
            self._first_cleanup_done = False
            self._final_cleanup_done = False

    def _shutdown_all_internal(self) -> None:
        """全量清理（内部版本，调用方须已持有 _cleanup_lock）。"""
        logger.info("[SW-Cleanup] 执行全量 SolidWorks 进程清理...")
        self.disconnect_sw_cached()
        self._terminate_sw_processes()
        logger.info("[SW-Cleanup] 全量清理完成")

    def do_first_cleanup(self) -> None:
        """首次 SW 全体清理（进入 SW 阶段前调用）。"""
        with self._cleanup_lock:
            if not self._first_cleanup_done:
                logger.info("[SW-Cleanup] === 首次全体 SW 进程清理（进入 SW 阶段前）===")
                self._shutdown_all_internal()
                self._first_cleanup_done = True

    def do_final_cleanup(self) -> None:
        """末次 SW 全体清理（SW 阶段全部完成后调用）。"""
        with self._cleanup_lock:
            if not self._final_cleanup_done:
                logger.info("[SW-Cleanup] === 末次全体 SW 进程清理（SW 阶段全部完成后）===")
                self._shutdown_all_internal()
                self._final_cleanup_done = True

    def reset_cleanup_state(self) -> None:
        """重置 SW 全量清理状态。"""
        with self._cleanup_lock:
            self._first_cleanup_done = False
            self._final_cleanup_done = False
            self.disconnect_sw_cached()
            logger.info("[SW-Cleanup] 已重置全量清理状态")

    # ------------------------------------------------------------------
    # SW 连接管理
    # ------------------------------------------------------------------

    def _connect_sw(self) -> Any:  # noqa: ANN401  COM 动态对象
        """三层降级连接/启动 SolidWorks 并返回 ISldWorks COM 对象。"""
        import win32com.client

        # ---- 第1层: 连接已运行的 SW ----
        logger.info("[SW-COM] 正在连接 SolidWorks (第1层: GetActiveObject)...")
        try:
            sw_app = win32com.client.GetActiveObject("SldWorks.Application")
            logger.info("[SW-COM] 已连接到运行中的 SolidWorks 实例")
            return sw_app
        except Exception as e1:
            logger.info(
                f"[SW-COM] GetActiveObject 失败 ({type(e1).__name__}: {e1})，"
                f"尝试启动新实例..."
            )

        # ---- 第2层: 通过 COM Dispatch 启动 ----
        logger.info("[SW-COM] 正在启动 SolidWorks (第2层: COM Dispatch)...")
        try:
            sw_app = win32com.client.Dispatch("SldWorks.Application")
            try:
                sw_app.UserControl = True
            except Exception:
                pass
            try:
                sw_app.Visible = bool(ENGINE_CONFIG.get("sw_visible", True))
            except Exception:
                pass
            logger.info("[SW-COM] SolidWorks 已通过 COM Dispatch 启动")
            time.sleep(OPERATION_TIMEOUTS["sw_dispatch_startup_delay"])
            return sw_app
        except Exception as e2:
            logger.warning(
                f"[SW-COM] COM Dispatch 失败 ({type(e2).__name__}: {e2})，"
                f"尝试备选方案..."
            )

        # ---- 第3层: 直接启动 exe ----
        logger.warning(
            "[SW-COM] 前两层连接均失败，将强制终止所有残留 SW 进程后重新启动。"
            "若您有其他 SolidWorks 窗口打开且包含未保存数据，请立即保存！"
        )
        logger.info("[SW-COM] 正在启动 SolidWorks (第3层: subprocess)...")
        self._terminate_sw_processes()
        if not self._launch_sw_process():
            logger.error("[SW-COM] 所有启动方式均失败，无法连接 SolidWorks")
            return None
        try:
            sw_app = win32com.client.GetActiveObject("SldWorks.Application")
        except Exception as e3:
            logger.error(f"[SW-COM] 第3层 GetActiveObject 失败: {e3}")
            return None
        try:
            sw_app.Visible = bool(ENGINE_CONFIG.get("sw_visible", True))
        except Exception:
            pass
        logger.info("[SW-COM] 已通过 subprocess 启动并连接 SolidWorks")
        return sw_app

    def _uninitialize_com_if_needed(self) -> None:
        """释放当前线程的 COM 初始化状态。"""
        if not self._com_initialized:
            return
        try:
            import pythoncom
            try:
                pythoncom.CoFreeUnusedLibraries()
            except Exception:
                pass
            pythoncom.CoUninitialize()
        except Exception as e:
            logger.debug(f"[SW-Cleanup] COM 反初始化异常: {e}")
        finally:
            self._com_initialized = False

    def _launch_sw_process(self) -> bool:
        """通过 subprocess 直接启动 SolidWorks.exe，轮询等待 COM 就绪。"""
        sw_exe = LOCAL_PATHS.get("sw_exe", "")
        if not sw_exe or not os.path.exists(sw_exe):
            logger.warning("[SW-COM] 未配置 SolidWorks 可执行文件路径 (sw_exe)，无法使用 subprocess 启动")
            return False

        logger.info(f"[SW-COM] 正在通过 subprocess 启动 SolidWorks: {sw_exe}")
        try:
            creation_flags = getattr(subprocess, 'CREATE_NO_WINDOW', 0) if os.name == "nt" else 0
            subprocess.Popen(
                [sw_exe],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                creationflags=creation_flags,
            )
            logger.info("[SW-COM] SolidWorks 进程已启动，等待 COM 接口就绪...")

            import win32com.client
            max_wait = 60
            for attempt in range(max_wait):
                time.sleep(1)
                try:
                    sw_app = win32com.client.GetActiveObject("SldWorks.Application")
                    if sw_app is not None:
                        logger.info(f"[SW-COM] SolidWorks COM 接口已就绪 (等待了 {attempt + 1} 秒)")
                        return True
                except (AttributeError, TypeError):
                    pass
                except Exception as e:
                    logger.debug(f"[SW-COM] 获取 SolidWorks COM 对象异常: {e}")
            logger.error(f"[SW-COM] 等待 SolidWorks 启动超时 ({max_wait} 秒)")
            return False
        except FileNotFoundError:
            logger.error(f"[SW-COM] 找不到 SolidWorks 可执行文件: {sw_exe}")
            return False
        except OSError as e:
            logger.error(f"[SW-COM] 启动 SolidWorks 进程失败: {e}")
            return False

    def _terminate_sw_processes(self):
        """清理可能残留的 SolidWorks 进程。"""
        if os.name != "nt":
            return
        try:
            if self._is_sw_process_running(timeout=10):
                logger.info("[SW-Cleanup] 检测到残留 SolidWorks 进程，正在终止...")
                kill_result = subprocess.run(
                    ["taskkill", "/f", "/im", "SLDWORKS.exe"],
                    capture_output=True, text=True, timeout=30,
                )
                if kill_result.returncode == 0:
                    logger.info("[SW-Cleanup] ✓ 残留 SolidWorks 进程已终止，等待 3 秒...")
                    time.sleep(3)
                else:
                    stderr = kill_result.stderr or ""
                    logger.warning(
                        f"[SW-Cleanup] taskkill 返回非零码 {kill_result.returncode}: "
                        f"{stderr.strip()}"
                    )
        except (subprocess.TimeoutExpired, OSError) as e:
            logger.warning(f"[SW-Cleanup] 检查/终止 SW 进程时异常: {e}")

    @staticmethod
    def _is_sw_process_running(timeout: int = 5) -> bool:
        """检查当前 Windows 会话中是否存在 SolidWorks 进程。"""
        if os.name != "nt":
            return False
        try:
            result = subprocess.run(
                ["tasklist", "/fi", "IMAGENAME eq SLDWORKS.exe", "/fo", "csv", "/nh"],
                capture_output=True, text=True, timeout=timeout,
            )
        except (subprocess.TimeoutExpired, OSError) as e:
            logger.debug(f"[SW-Cleanup] 检查 SolidWorks 进程时异常: {e}")
            return False
        stdout = result.stdout or ""
        return "SLDWORKS.exe" in stdout

    @staticmethod
    def _is_com_rpc_failure(exc: Exception | None) -> bool:
        """判断 COM 异常是否表示 SolidWorks RPC/进程崩溃。"""
        if exc is None:
            return False
        text = str(exc)
        return any(
            marker in text
            for marker in (
                "-2147023170",
                "0x800706BE",
                "-2147023174",
                "0x800706BA",
                "远程过程调用失败",
                "RPC_S_CALL_FAILED",
                "RPC_S_SERVER_UNAVAILABLE",
            )
        )

    def _recover_after_open_rpc_failure(
        self,
        sw_model: str,
        doc_type: int,
    ) -> tuple[bool, Any | None, Any | None]:
        """OpenDoc6 RPC 失败时清理残留进程并重连一次。"""
        if not self._is_com_rpc_failure(self._last_open_error):
            return False, None, None

        logger.warning(
            "[SW-COM] OpenDoc6 检测到 COM/RPC 失败，"
            "将终止残留 SolidWorks 进程并重试一次"
        )
        self._terminate_sw_processes()
        self._uninitialize_com_if_needed()
        try:
            import pythoncom

            pythoncom.CoInitialize()
            self._com_initialized = True
        except Exception as e:
            logger.debug(f"[SW-COM] RPC 恢复时 COM 初始化异常: {e}")

        sw_app = self._connect_sw()
        if sw_app is None:
            return True, None, None
        self._hide_document_window_for_recovery(sw_app, doc_type)
        return True, sw_app, self._open_sw_model(sw_app, sw_model, doc_type)

    @staticmethod
    def _hide_document_window_for_recovery(
        sw_app: Any,
        doc_type: int,
    ) -> None:  # noqa: ANN401  COM 动态对象
        """恢复重试时隐藏文档窗口，避开崩溃的 UI 渲染路径。"""
        try:
            sw_app.Visible = False
        except Exception as e:
            logger.debug(f"[SW-COM] 设置恢复打开隐藏主窗口失败: {e}")
        try:
            sw_app.DocumentVisible(False, doc_type)
        except Exception as e:
            logger.debug(f"[SW-COM] 设置恢复打开隐藏窗口失败: {e}")

    def _open_sw_model(self, sw_app: Any, sw_model: str, doc_type: int) -> Any:  # noqa: ANN401  COM 动态对象
        """通过 OpenDoc6 打开 SW 模型文件并验证 COM 代理有效性。"""
        import win32com.client
        import pythoncom

        self._last_open_error = None
        logger.info(f"[SW-COM] 正在打开模型 (OpenDoc6): {os.path.basename(sw_model)}")
        open_errors = win32com.client.VARIANT(pythoncom.VT_BYREF | pythoncom.VT_I4, 0)
        open_warnings = win32com.client.VARIANT(pythoncom.VT_BYREF | pythoncom.VT_I4, 0)

        try:
            doc = sw_app.OpenDoc6(
                sw_model,
                doc_type,
                self._SW_OPEN_SILENT | self._SW_OPEN_READONLY,
                "",
                open_errors,
                open_warnings,
            )
            logger.info(
                f"[SW-COM] OpenDoc6: Errors={open_errors.value}, "
                f"Warnings={open_warnings.value}"
            )
            if open_errors.value != 0:
                logger.warning(
                    f"[SW-COM] OpenDoc6 返回错误码 {open_errors.value}，"
                    f"模型可能存在问题（缺失参考/重建错误）"
                )
        except Exception as open_err:
            self._last_open_error = open_err
            logger.error(
                f"[SW-COM] OpenDoc6 异常 ({type(open_err).__name__}: {open_err})"
            )
            return None

        if doc is None:
            logger.error(
                f"[SW-COM] 无法打开 SW 模型: {sw_model}"
                f"（文件可能损坏、版本不兼容，或路径含特殊字符）"
            )
            return None

        if not self._verify_com_object(doc, "IModelDoc2"):
            logger.error(
                "[SW-COM] OpenDoc6 返回了无效的文档 COM 代理，"
                "模型可能未正确加载"
            )
            try:
                sw_app.CloseDoc(os.path.basename(sw_model))
            except Exception:
                pass
            return None

        logger.info(f"[SW-COM] 模型已打开: {os.path.basename(sw_model)}")
        return doc

    @staticmethod
    def _capture_design_table_snapshot() -> tuple[str, bytes] | None:
        """读取外部设计表快照，防止 SolidWorks 回写污染源参数表。"""
        excel_path = str(LOCAL_PATHS.get("excel") or "")
        if not excel_path or not os.path.exists(excel_path):
            return None
        try:
            with open(excel_path, "rb") as f:
                return excel_path, f.read()
        except OSError as e:
            logger.warning(f"[SW-DesignTable] 无法读取源 Excel 快照: {excel_path}: {e}")
            return None

    def _restore_design_table_snapshot(self, snapshot: tuple[str, bytes] | None) -> bool:
        """如果 SW 修改了外部设计表，立即恢复导出前的源 Excel。"""
        if snapshot is None:
            return True
        excel_path, original_bytes = snapshot
        try:
            current_bytes = b""
            if os.path.exists(excel_path):
                with open(excel_path, "rb") as f:
                    current_bytes = f.read()
            if current_bytes == original_bytes:
                return True
            with open(excel_path, "wb") as f:
                f.write(original_bytes)
            logger.warning(
                "[SW-DesignTable] SolidWorks 修改了源 Excel 参数表，"
                f"已恢复原始内容: {excel_path}"
            )
            return True
        except OSError as e:
            self.last_error = f"[SW-DesignTable] 恢复源 Excel 参数表失败: {excel_path}: {e}"
            logger.error(self.last_error)
            return False

    def _disconnect_sw(self, sw_app, doc, sw_model: str):
        """清理 SW COM 资源：关闭文档 → 释放 COM。

        SolidWorks 进程退出统一由全量清理流程处理。
        """
        import pythoncom

        # 步骤 1: 关闭模型文档
        if doc is not None and ENGINE_CONFIG.get("sw_close_doc_on_finish", True):
            try:
                title = doc.GetTitle()
            except Exception:
                title = os.path.basename(sw_model)
            try:
                sw_app.CloseDoc(title)
                logger.info(f"[SW-Cleanup] 已关闭模型文档: {title}")
            except Exception as e_doc:
                logger.debug(f"[SW-Cleanup] 关闭模型文档异常: {e_doc}")

        # 步骤 2: 释放 COM 资源
        del doc
        del sw_app
        gc.collect()
        for _ in range(2):
            try:
                pythoncom.CoFreeUnusedLibraries()
            except Exception:
                pass
            time.sleep(0.5)
        pythoncom.CoUninitialize()

    # ------------------------------------------------------------------
    # 已链接设计表格式预验证
    # ------------------------------------------------------------------

    def _validate_design_table(self, excel_path: str) -> list[str]:
        """预验证 Excel 设计表格式，返回诊断警告列表。"""
        warnings = []
        logger.info(f"[SW-DesignTable] 预验证 Excel 设计表格式: {os.path.basename(excel_path)}")

        if not os.path.exists(excel_path):
            warnings.append(f"Excel 文件不存在: {excel_path}")
            logger.warning(f"[SW-DesignTable] {warnings[-1]}")
            return warnings

        file_size = os.path.getsize(excel_path)
        logger.info(f"[SW-DesignTable] 文件大小: {file_size} bytes")

        try:
            import openpyxl
            wb = None
            try:
                wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
                ws = wb.active

                try:
                    row1_cells = list(ws.iter_rows(min_row=1, max_row=1))
                    if not row1_cells:
                        warnings.append("Excel 第1行缺失（空表格）")
                        logger.warning(f"[SW-DesignTable] {warnings[-1]}")
                        return warnings
                    row1 = [cell.value for cell in row1_cells[0]]
                except (IndexError, StopIteration):
                    warnings.append("Excel 第1行缺失（空表格）")
                    logger.warning(f"[SW-DesignTable] {warnings[-1]}")
                    return warnings
                row1_text = " ".join(str(v) for v in row1 if v is not None)
                logger.info(f"[SW-DesignTable] 第1行内容: {row1_text[:120]}")

                has_design_table_header = "Design Table" in row1_text or "设计表" in row1_text
                if not has_design_table_header:
                    warnings.append(
                        f"Excel 第1行缺少 SW 设计表头（应包含 'Design Table'），"
                        f"实际内容: {row1_text[:80]}"
                    )
                    logger.warning(f"[SW-DesignTable] {warnings[-1]}")

                try:
                    row2_cells = list(ws.iter_rows(min_row=2, max_row=2))
                    if row2_cells:
                        row2 = [cell.value for cell in row2_cells[0]]
                        row2_text = " | ".join(str(v) for v in row2 if v is not None)
                        logger.info(f"[SW-DesignTable] 第2行内容: {row2_text[:200]}")
                        has_param_headers = any(
                            isinstance(v, str) and ("$" in v or "@" in v)
                            for v in row2 if v is not None
                        )
                        if not has_param_headers:
                            warnings.append(
                                f"Excel 第2行似乎不含 SW 参数列头（期望格式如 '$PRP@Dimension1'），"
                                f"实际内容: {row2_text[:100]}"
                            )
                            logger.warning(f"[SW-DesignTable] {warnings[-1]}")
                except (IndexError, StopIteration):
                    warnings.append("Excel 第2行缺失（应为参数列头行）")
                    logger.warning(f"[SW-DesignTable] {warnings[-1]}")

                data_rows = 0
                for row in ws.iter_rows(min_row=3, values_only=True):
                    if row[0] is not None:
                        data_rows += 1
                logger.info(f"[SW-DesignTable] 数据行数 (第3行起): {data_rows}")

            finally:
                if wb is not None:
                    wb.close()

        except ImportError:
            logger.warning("[SW-DesignTable] openpyxl 未安装，跳过 Excel 格式预验证")
        except Exception as e:
            warnings.append(f"Excel 格式预验证异常: {type(e).__name__}: {e}")
            logger.warning(f"[SW-DesignTable] {warnings[-1]}", exc_info=True)

        if not warnings:
            logger.info("[SW-DesignTable] Excel 设计表格式预验证通过")
        return warnings

    # ------------------------------------------------------------------
    # STEP 导出
    # ------------------------------------------------------------------

    @staticmethod
    def _guess_sw_doc_type(path: str) -> int:
        """根据文件扩展名猜测 SW 文档类型。"""
        ext = os.path.splitext(path)[1].lower()
        if ext == ".sldasm":
            return SWExecutor._SW_DOC_ASSEMBLY
        return SWExecutor._SW_DOC_PART

    @staticmethod
    def _verify_com_object(obj, label: str = "COM对象") -> bool:
        """验证 COM 对象是否有效（非 None 且代理仍存活）。"""
        if obj is None:
            logger.error(f"[SW-COM] {label} 为 None")
            return False

        for method_name in ("GetTitle", "GetPathName", "GetType"):
            try:
                val = getattr(obj, method_name)
                if callable(val):
                    _ = val()
                logger.debug(f"[SW-COM] {label} 有效 (通过 {method_name})")
                return True
            except AttributeError:
                continue
            except Exception as e:
                logger.debug(
                    f"[SW-COM] {label} {method_name} 失败: "
                    f"{type(e).__name__}"
                )
                continue

        logger.warning(
            f"[SW-COM] {label} 所有验证方法均失败，"
            f"对象可能为无效 COM 代理"
        )
        return False

    def _verify_step_exports(self, step_dir: str) -> int:
        """安全网校验：扫描所有构型的 STEP 输出文件，补标记状态数据库。"""
        all_configs = self.state.get_all_configs()
        missing_configs: list[int] = []
        found_configs: list[int] = []
        already_completed: list[int] = []

        logger.info("[SW-Export] STEP 导出完毕，正在校验各构型 STEP 文件...")
        logger.info(f"[SW-Export] 输出目录: {step_dir}")

        for cn in all_configs:
            filename = get_step_filename("sw", cn)
            if not filename:
                logger.warning(f"[SW-Export] 构型{cn}: 无法生成 STEP 文件名，跳过校验")
                continue
            expected_file = os.path.join(step_dir, filename)
            current_status = self.state.get_step_status(cn, "sw")
            if os.path.exists(expected_file):
                if current_status == STATUS_COMPLETED:
                    already_completed.append(cn)
                    logger.debug(f"[SW-Export] 构型{cn} (文件监控器已标记)")
                else:
                    self.state.set_step_status(cn, "sw", STATUS_COMPLETED)
                    found_configs.append(cn)
                    logger.debug(f"[SW-Export] 构型{cn} (安全网补标记)")
            else:
                self.state.set_step_status(
                    cn, "sw", STATUS_ERROR,
                    f"STEP 导出完毕但文件缺失: {filename}"
                )
                missing_configs.append(cn)
                logger.warning(
                    f"[SW-Export] 构型{cn} STEP 缺失 "
                    "— 可能原因: 构型重建失败 / 设计表参数错误"
                )

        total_found = len(already_completed) + len(found_configs)
        logger.info(
            f"[SW-Export] STEP 校验完成: "
            f"{total_found}/{len(all_configs)} 成功"
            f"（文件监控器实时: {len(already_completed)}，安全网: {len(found_configs)}）"
        )
        return total_found
