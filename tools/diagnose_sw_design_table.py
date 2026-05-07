"""
===============================================================================
SolidWorks 设计表诊断工具
诊断 Excel 参数表格式、SW 模型参数结构和 InsertFamilyTableOpen 行为。

用法: python tools/diagnose_sw_design_table.py
===============================================================================
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from engine.config import LOCAL_PATHS

EXCEL_PATH = LOCAL_PATHS["excel"]
SW_MODEL = LOCAL_PATHS["sw_model"]


def diagnose_excel():
    """诊断 Excel 参数表格式。"""
    print("=" * 60)
    print("1) Excel 参数表诊断")
    print("=" * 60)

    import openpyxl
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active
    print(f"  文件: {EXCEL_PATH}")
    print(f"  工作表: {ws.title}")

    # 打印前5行
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), 1):
        values = [str(c)[:40] if c is not None else "(空)" for c in row[:8]]
        print(f"  行{row_idx}: {' | '.join(values)}")

    # 统计行数
    data_rows = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is not None:
            data_rows += 1
    print(f"  数据行数 (第3行起): {data_rows}")
    wb.close()


def diagnose_sw_model():
    """诊断 SW 模型参数结构和配置信息。"""
    print()
    print("=" * 60)
    print("2) SW 模型结构诊断")
    print("=" * 60)

    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
    except ImportError:
        print("  pywin32 未安装，跳过")
        return

    try:
        # 连接 SW
        sw_app = None
        try:
            sw_app = win32com.client.GetActiveObject("SldWorks.Application")
            print("  已连接运行中的 SW")
        except Exception:
            sw_app = win32com.client.Dispatch("SldWorks.Application")
            sw_app.Visible = True
            print("  已启动新 SW 实例")
            import time
            time.sleep(8)

        if sw_app is None:
            print("  无法连接 SW")
            return

        # SW version
        try:
            rev = sw_app.RevisionNumber
            print(f"  SW 版本: {rev}")
        except Exception:
            pass

        # 打开模型
        print(f"  正在打开模型: {os.path.basename(SW_MODEL)}")
        open_errors = win32com.client.VARIANT(pythoncom.VT_BYREF | pythoncom.VT_I4, 0)
        open_warnings = win32com.client.VARIANT(pythoncom.VT_BYREF | pythoncom.VT_I4, 0)
        doc_type = 1  # swDocPART

        doc = sw_app.OpenDoc6(SW_MODEL, doc_type, 1, "", open_errors, open_warnings)
        print(f"  OpenDoc6: Errors={open_errors.value}, Warnings={open_warnings.value}")

        if doc is None:
            print("  无法打开模型")
            return

        # 获取文档标题
        try:
            title = doc.GetTitle()
            print(f"  文档标题: {title}")
        except Exception as e:
            print(f"  GetTitle 失败: {e}")

        # 获取配置列表
        print()
        print("  配置列表:")
        config_names = []
        try:
            raw = doc.GetConfigurationNames()
            if isinstance(raw, (tuple, list)):
                config_names = list(raw)
            else:
                config_names = [str(raw)]
            for i, cfg in enumerate(config_names[:10]):
                print(f"    [{i}] {cfg}")
            if len(config_names) > 10:
                print(f"    ... 共 {len(config_names)} 个配置")
        except Exception as e:
            print(f"    GetConfigurationNames 失败: {e}")

        # 获取参数列表
        print()
        print("  模型参数 (全部):")
        model_param_names = []
        try:
            all_params = doc.Extension.GetParameters()
            if all_params:
                for p in all_params:
                    try:
                        name = p.Name
                        sv = p.SystemValue
                        val = p.Value
                        print(f"    {name:35s}  Value={val}  SystemValue={sv}")
                        model_param_names.append(name)
                    except Exception:
                        pass
                print(f"    共 {len(model_param_names)} 个参数")
            else:
                print("    GetParameters 返回空")
        except Exception as e:
            print(f"    GetParameters 异常: {type(e).__name__}: {e}")

        # 检查设计表
        print()
        print("  设计表检查:")
        try:
            dt = doc.GetDesignTable()
            if dt is not None:
                print(f"    ✓ 模型已有内嵌设计表")
                try:
                    dt_name = dt.Name
                    print(f"      设计表名称: {dt_name}")
                except Exception:
                    pass
            else:
                print(f"    模型暂无设计表")
        except Exception as e:
            print(f"    GetDesignTable 异常: {e}")

        # 测试 InsertFamilyTableOpen
        print()
        print("  测试 InsertFamilyTableOpen:")
        try:
            result = doc.InsertFamilyTableOpen(EXCEL_PATH)
            print(f"    返回: {result}")
            if result:
                print(f"    ✓ InsertFamilyTableOpen 成功!")
                dt2 = doc.GetDesignTable()
                if dt2 is not None:
                    print(f"    ✓ GetDesignTable 返回有效对象")
            else:
                print(f"    ✗ 返回 False — Excel 格式可能不兼容")
        except Exception as e:
            print(f"    COM 异常: {type(e).__name__}: {e}")
            # Try to understand the error
            if hasattr(e, 'args') and e.args:
                print(f"    错误详情: {e.args}")

        # Test InsertFamilyTable (no-arg version)
        print()
        print("  测试 InsertFamilyTable() (无参数版):")
        try:
            # First delete any existing one
            try:
                doc.DeleteDesignTable()
                print("    已删除已有设计表")
            except Exception:
                pass

            doc.InsertFamilyTable()
            print("    ✓ InsertFamilyTable() 成功")
            dt3 = doc.GetDesignTable()
            if dt3 is not None:
                print("    ✓ GetDesignTable 返回有效对象")
                # Try to set up design table entries
                try:
                    dt3.Updatable = False
                    print("    已设置 Updatable=False")
                except Exception:
                    pass
        except Exception as e:
            print(f"    COM 异常: {type(e).__name__}: {e}")

        # Try direct parameter setting approach
        print()
        print("  ================================================================")
        print("  备选方案: 直接通过 COM API 设置参数 (绕过设计表)")
        print("  ================================================================")
        print(f"  参数列表 (含 SystemValue 和 Expression):")
        try:
            all_params = doc.Extension.GetParameters()
            for p in all_params[:15]:
                try:
                    name = p.Name
                    try:
                        sv = p.SystemValue
                    except Exception:
                        sv = "N/A"
                    try:
                        expr = p.Expression
                    except Exception:
                        expr = "N/A"
                    print(f"    {name:30s}  SystemValue={sv}  Expression={expr}")
                except Exception:
                    pass
        except Exception:
            pass

        # Try activating each configuration and listing their parameter values
        print()
        print("  各配置参数值对比 (前5个配置):")
        try:
            config_names = doc.GetConfigurationNames()
            for cfg_name in (config_names[:5] if config_names else []):
                print(f"\n    配置: {cfg_name}")
                doc.ShowConfiguration2(cfg_name)
                try:
                    all_params2 = doc.Extension.GetParameters()
                    for p in all_params2[:10]:
                        try:
                            name = p.Name
                            sv = p.SystemValue
                            print(f"      {name:30s} = {sv}")
                        except Exception:
                            pass
                except Exception as e_p:
                    print(f"      获取参数失败: {e_p}")
        except Exception as e:
            print(f"    配置切换异常: {e}")

        # 关闭文档
        try:
            sw_app.CloseDoc(os.path.basename(SW_MODEL))
            print(f"\n  已关闭文档")
        except Exception:
            pass

    finally:
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    print("SolidWorks 设计表诊断工具")
    print("=" * 60)
    print()

    # 检查文件
    if not os.path.exists(EXCEL_PATH):
        print(f"✗ Excel 文件不存在: {EXCEL_PATH}")
    else:
        diagnose_excel()

    if not os.path.exists(SW_MODEL):
        print(f"✗ SW 模型文件不存在: {SW_MODEL}")
    else:
        diagnose_sw_model()

    print()
    print("=" * 60)
    print("诊断完成")
    print("=" * 60)
