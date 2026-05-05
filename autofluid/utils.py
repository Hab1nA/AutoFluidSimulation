"""通用工具函数。"""

from __future__ import annotations

import subprocess
import time
from pathlib import Path
from typing import Iterable, List, Optional

from openpyxl import load_workbook

from .models import ConfigRow


def normalize_value(value: object) -> Optional[float]:
    """将 Excel 单元格数据转换为浮点或空值。"""

    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_excel_configs(excel_path: str) -> List[ConfigRow]:
    """读取 Excel，按需求从第 3 行开始获取构型列表。"""

    workbook = load_workbook(excel_path, data_only=True)
    sheet = workbook.active
    configs: List[ConfigRow] = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        name = row[0]
        if name is None or str(name).strip() == "":
            break
        params = [normalize_value(value) for value in row[1:5]]
        while len(params) < 4:
            params.append(None)
        configs.append(
            ConfigRow(
                name=str(name).strip(),
                param1=params[0],
                param2=params[1],
                param3=params[2],
                param4=params[3],
            )
        )
    return configs


def ensure_dir(path: str) -> None:
    """确保目录存在。"""

    Path(path).mkdir(parents=True, exist_ok=True)


def run_subprocess(command: List[str], timeout: Optional[int] = None) -> None:
    """运行子进程并在失败时抛出异常。"""

    subprocess.run(command, check=True, timeout=timeout)


def find_step_file(step_dir: str, config_name: str, extensions: Iterable[str]) -> Optional[Path]:
    """查找构型对应的 STEP 文件。"""

    directory = Path(step_dir)
    for ext in extensions:
        matches = sorted(directory.glob(f"{config_name}*.{ext}"), key=lambda item: item.name)
        if matches:
            return matches[0]
    return None


def now_iso() -> str:
    """获取当前时间字符串。"""

    return time.strftime("%Y-%m-%d %H:%M:%S")
